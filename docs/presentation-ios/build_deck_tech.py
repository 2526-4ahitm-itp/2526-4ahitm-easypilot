#!/usr/bin/env python3
"""Builds the EasyPilot iOS technology deep-dive deck.

One slide per technology: a minimal, readable code snippet plus a bullet-style
explanation. Swift and the 10 Hz communicator are intentionally omitted.
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
LOGOS = os.path.join(HERE, "logos")

# ---- Theme ----
BG     = RGBColor(0x0A, 0x0F, 0x1E)
CARD   = RGBColor(0x14, 0x1C, 0x33)
ACCENT = RGBColor(0x29, 0x79, 0xFF)
TEXT   = RGBColor(0xE9, 0xEE, 0xFB)
MUTED  = RGBColor(0x8A, 0x97, 0xB8)
CODEBG = RGBColor(0x0C, 0x13, 0x26)
GREEN  = RGBColor(0x33, 0xD9, 0x80)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]

# ---- Syntax highlighting colours ----
KW = RGBColor(0xFF, 0x7A, 0xB2)  # keywords / punctuation
TY = RGBColor(0x7E, 0xE0, 0xC8)  # types
ST = RGBColor(0xC3, 0xE8, 0x8D)  # strings
CM = RGBColor(0x5D, 0x6B, 0x8F)  # comments
NU = RGBColor(0xF7, 0xC8, 0x73)  # numbers
FN = RGBColor(0x82, 0xAA, 0xFF)  # function names
PL = TEXT                         # plain


def slide():
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def textbox(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=6):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for line in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(sp_after)
        p.space_before = Pt(0)
        for txt, size, color, bold, fname in line:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Consolas" if fname == "mono" else "Helvetica Neue"
    return tb


def kicker(s, txt):
    textbox(s, Inches(0.85), Inches(0.55), Inches(11), Inches(0.5),
            [[(txt.upper(), 15, ACCENT, True, None)]])


def title(s, txt, color=TEXT):
    textbox(s, Inches(0.85), Inches(0.95), Inches(11.6), Inches(1.1),
            [[(txt, 40, color, True, None)]])


def card(s, x, y, w, h, border=ACCENT):
    c = s.shapes.add_shape(5, x, y, w, h)
    c.fill.solid(); c.fill.fore_color.rgb = CARD
    c.line.color.rgb = border; c.line.width = Pt(1.25)
    c.shadow.inherit = False
    try:
        c.adjustments[0] = 0.06
    except Exception:
        pass
    return c


def fit(path, max_w_in, max_h_in):
    im = Image.open(path)
    iw, ih = im.size
    ar = iw / ih
    w, h = max_w_in, max_w_in / ar
    if h > max_h_in:
        h = max_h_in
        w = max_h_in * ar
    return Inches(w), Inches(h)


def add_code_box(s, x, y, w, h, code_lines):
    box = s.shapes.add_shape(5, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG
    box.line.color.rgb = ACCENT; box.line.width = Pt(1)
    box.shadow.inherit = False
    try:
        box.adjustments[0] = 0.03
    except Exception:
        pass
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.35)
    tf.margin_top = Inches(0.25)
    tf.margin_right = Inches(0.2)
    first = True
    for segs in code_lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(2)
        p.line_spacing = 1.0
        for txt, col in segs:
            r = p.add_run()
            r.text = txt
            r.font.name = "Consolas"
            r.font.size = Pt(13.5)
            r.font.color.rgb = col


def add_bullets(s, x, y, w, h, bullets):
    box = s.shapes.add_shape(5, x, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = CARD
    box.line.color.rgb = ACCENT; box.line.width = Pt(1)
    box.shadow.inherit = False
    try:
        box.adjustments[0] = 0.03
    except Exception:
        pass
    tb = textbox(s, x + Inches(0.35), y + Inches(0.25), w - Inches(0.55), h - Inches(0.45),
                 [[("✦  " + b, 15, TEXT, False, None)] for b in bullets], sp_after=10)
    return tb


def tech_slide(logo, tech_name, subtitle, code_lines, bullets):
    s = slide()
    kicker(s, "Technology Deep-Dive")

    # logo + title row
    w, h = fit(os.path.join(LOGOS, logo), 1.0, 1.0)
    s.shapes.add_picture(os.path.join(LOGOS, logo), Inches(0.85), Inches(0.92), width=w, height=h)

    textbox(s, Inches(2.05), Inches(0.95), Inches(10.4), Inches(0.7),
            [[(tech_name, 34, TEXT, True, None)]])
    textbox(s, Inches(2.05), Inches(1.55), Inches(10.4), Inches(0.5),
            [[(subtitle, 16, MUTED, False, None)]])

    # two-column layout
    add_code_box(s, Inches(0.85), Inches(2.35), Inches(5.9), Inches(4.4), code_lines)
    add_bullets(s, Inches(7.0), Inches(2.35), Inches(5.45), Inches(4.4), bullets)
    return s


# ===========================================================================
# Slide 1 — Title
# ===========================================================================
s = slide()
bar = s.shapes.add_shape(1, 0, 0, Inches(0.22), SH)
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
bar.shadow.inherit = False
kicker(s, "EasyPilot · 4AHITM Drone Project")
title(s, "Technologie-Überblick")
textbox(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.4),
        [[("Eine Folie pro Technologie: ein minimales Code-Beispiel", 22, MUTED, False, None)],
         [("plus die wichtigsten Eigenschaften in Stichworten.", 22, MUTED, False, None)]],
        sp_after=6)
textbox(s, Inches(0.9), Inches(4.2), Inches(11.5), Inches(0.6),
        [[("Ohne Swift und ohne den 10 Hz Communicator.", 17, GREEN, True, None)]])

# ===========================================================================
# Slide 2 — SwiftUI
# ===========================================================================
tech_slide(
    "swiftui.png", "SwiftUI",
    "Declaratives UI-Framework von Apple",
    [
        [("struct ", KW), ("GreetingView", TY), (": ", PL), ("View", TY), (" {", PL)],
        [("    @", KW), ("State", TY), (" var", KW), (" name = ", PL), ("\"World\"", ST)],
        [("    var", KW), (" body: some ", PL), ("View", TY), (" {", PL)],
        [("        VStack", TY), (" {", PL)],
        [("            Text", TY), ("(\"Hello, \\(name)!\"", ST), (")", PL)],
        [("            Button", TY), ("(", PL), ("\"Tap me\"", ST), (") {", PL)],
        [("                name = ", PL), ("\"EasyPilot\"", ST)],
        [("            }", PL)],
        [("        }", PL)],
        [("    }", PL)],
        [("}", PL)],
    ],
    [
        "Declarative: Beschreibung WAS angezeigt wird",
        "@State für lokale UI-Zustände",
        "Live-Updates bei Datenänderung",
        "Kombiniert Views durch Stacks",
        "Native Apple-Framework",
    ]
)

# ===========================================================================
# Slide 3 — Network.framework
# ===========================================================================
tech_slide(
    "network.png", "Network.framework",
    "Low-Level Networking in iOS",
    [
        [("let", KW), (" endpoint = ", PL), ("NWEndpoint", TY), (".hostPort(", PL),
         ("host: ", KW), ("\"192.168.4.1\"", ST), (", port: ", PL), ("80", NU), (")", PL)],
        [("let", KW), (" parameters = ", PL), ("NWParameters", TY), (".tcp", PL)],
        [("let", KW), (" connection = ", PL), ("NWConnection", TY),
         ("(to: endpoint, using: parameters)", PL)],
        [("connection.", PL), ("start", FN), ("(queue: .global())", PL)],
        [("connection.", PL), ("send", FN), ("(content: data, completion: .contentProcessed { _ in", PL),
         ("", PL)],
        [("    print", FN), ("(\"sent\"", ST), (")", PL)],
        [("})", PL)],
    ],
    [
        "TCP / UDP / TLS auf einer Ebene",
        "NWListener für eingehende Verbindungen",
        "Wird für UDP-Discovery genutzt",
        "Warteschlangen-basiert (queues)",
        "Ersatz für ältere Socket-APIs",
    ]
)

# ===========================================================================
# Slide 4 — SceneKit
# ===========================================================================
tech_slide(
    "scenekit.png", "SceneKit",
    "3D-Rendering-Engine für iOS",
    [
        [("let", KW), (" scene = ", PL), ("SCNScene", TY), ("(named: ", PL), ("\"drone.usdz\"", ST), (")!", PL)],
        [("let", KW), (" drone = scene.", PL), ("rootNode", FN),
         (".childNode(withName: ", PL), ("\"drone\"", ST), (", recursively: ", PL), ("true", KW), (")", PL)],
        [("drone?.", PL), ("runAction", FN), ("(", PL)],
        [("    .", PL), ("rotateBy", FN), ("(x: ", PL), ("0", NU), (", y: ", PL), ("1", NU), (", z: ", PL), ("0", NU), (", duration: ", PL), ("2", NU), (")", PL)],
        [(")", PL)],
    ],
    [
        "Apples High-Level 3D-Framework",
        "Lädt .usdz Modelle",
        "Szenengraph aus SCNNodes",
        "Animationen & Actions",
        "Physik-Engine integrierbar",
    ]
)

# ===========================================================================
# Slide 5 — CoreMotion
# ===========================================================================
tech_slide(
    "coremotion.png", "CoreMotion",
    "Zugriff auf Bewegungssensoren",
    [
        [("let", KW), (" manager = ", PL), ("CMMotionManager", TY), ("()", PL)],
        [("manager.", PL), ("deviceMotionUpdateInterval", FN), (" = ", PL), ("1.0 / 10.0", NU), (" // 10 Hz", CM)],
        [("manager.", PL), ("startDeviceMotionUpdates", FN),
         ("(to: .main) { data, error ", PL), ("in", KW), ("", PL)],
        [("    guard", KW), (" let", KW), (" attitude = data?.", PL), ("attitude", FN), (" else", KW), (" { return }", PL)],
        [("    let", KW), (" pitchDeg = attitude.", PL), ("pitch", FN), (" * ", PL), ("180", NU), (" / .pi", PL)],
        [("    print", FN), ("(\"Pitch: \\(pitchDeg)°\"", ST), (")", PL)],
        [("}", PL)],
    ],
    [
        "Gyroskop & Accelerometer",
        "DeviceMotion fusioniert Sensoren",
        "Liefert pitch / roll / yaw",
        "Konfigurierbare Update-Rate",
        "Nebenläufig via OperationQueue",
    ]
)

# ===========================================================================
# Slide 6 — WebSocket
# ===========================================================================
tech_slide(
    "websocket.png", "WebSocket",
    "Vollduplex-Kommunikation über TCP",
    [
        [("let", KW), (" url = ", PL), ("URL", TY), ("(string: ", PL), ("\"ws://drone.local/ws\"", ST), (")!", PL)],
        [("let", KW), (" task = ", PL), ("URLSession", TY), (".shared.", PL), ("webSocketTask", FN), ("(with: url)", PL)],
        [("task.", PL), ("resume", FN), ("()", PL)],
        [("task.", PL), ("send", FN), ("(", PL), (".string", TY), ("(", PL), ("\"{\"cmd\":\"ARM\"}\"", ST), (") { error ", PL), ("in", KW)],
        [("    if", KW), (" error == ", PL), ("nil", KW), (" { print", FN), ("(\"armed\"", ST), (") }", PL)],
        [("}", PL)],
    ],
    [
        "Vollduplex: gleichzeitig senden & empfangen",
        "Persistente Verbindung",
        "Geringer Overhead pro Nachricht",
        "Echtzeit-Telemetrie geeignet",
        "Bidirektionale JSON-Kommandos",
    ]
)

# ===========================================================================
# Slide 7 — ESP32 / Arduino
# ===========================================================================
tech_slide(
    "esp.png", "ESP32 / Arduino",
    "On-drone Mikrocontroller-Firmware",
    [
        [("#include <WiFi.h>", PL)],
        [("#include <WebSocketsServer.h>", PL)],
        [("WebSocketsServer ", KW), ("webSocket", PL), (" = ", PL), ("WebSocketsServer(81)", PL), (";", PL)],
        [("void ", KW), ("setup", FN), ("() {", PL)],
        [("    WiFi.", PL), ("softAP", FN), ("(\"EasyPilot\"", ST), (");", PL)],
        [("    webSocket.", PL), ("begin", FN), ("();", PL)],
        [("}", PL)],
        [("void ", KW), ("loop", FN), ("() {", PL)],
        [("    webSocket.", PL), ("loop", FN), ("();", PL), (" // auf Nachrichten prüfen", CM)],
        [("}", PL)],
    ],
    [
        "Dual-Core Mikrocontroller",
        "Integriertes WiFi & Bluetooth",
        "Programmiert mit Arduino-Framework",
        "Läuft als Access-Point",
        "WebSocket-Server für Befehle",
    ]
)

# ===========================================================================
# Slide 8 — Wrap-up
# ===========================================================================
s = slide()
kicker(s, "Zusammenfassung")
title(s, "Was wir gesehen haben")
card(s, Inches(0.85), Inches(2.1), Inches(11.6), Inches(3.9))
bullets = [
    "SwiftUI: deklarative UI mit Live-Bindings",
    "Network.framework: Low-Level Netzwerk-API",
    "SceneKit: 3D-Rendering und Simulator",
    "CoreMotion: Sensorfusion für Neigungsdaten",
    "WebSocket: Echtzeit-Kommunikation mit der Drohne",
    "ESP32: Onboard Mikrocontroller & WebSocket-Server",
]
textbox(s, Inches(1.15), Inches(2.45), Inches(11), Inches(3.4),
        [[("•  " + b, 18, TEXT, False, None)] for b in bullets], sp_after=12)
textbox(s, Inches(0.85), Inches(6.4), Inches(11.6), Inches(0.6),
        [[("EasyPilot · 4AHITM HTL · Technologie-Überblick", 15, MUTED, False, None)]],
        align=PP_ALIGN.CENTER)

out = os.path.join(HERE, "EasyPilot-iOS-Technologies.pptx")
prs.save(out)
print("Saved:", out, "—", len(prs.slides._sldIdLst), "slides")
